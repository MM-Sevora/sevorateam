"""
Feature Document Parser API
Parses DOCX feature documents and generates Epics, User Stories, and Tasks using AI
"""
import os
import uuid
import base64
import logging
from datetime import datetime, timezone
from typing import List, Optional, Dict, Any
from io import BytesIO

from fastapi import APIRouter, HTTPException, UploadFile, File, Depends, Form
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from pydantic import BaseModel
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import jwt

# Database connection
from motor.motor_asyncio import AsyncIOMotorClient
client = AsyncIOMotorClient(os.environ.get('MONGO_URL', 'mongodb://localhost:27017'))
db = client[os.environ.get('DB_NAME', 'sevora_production')]

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/feature-parser", tags=["Feature Parser"])
security = HTTPBearer()


# ============== Auth Helper ==============

async def get_current_user(credentials: HTTPAuthorizationCredentials = Depends(security)):
    """Verify JWT token and return current user"""
    try:
        token = credentials.credentials
        payload = jwt.decode(token, os.environ.get('JWT_SECRET', 'fallback-secret'), algorithms=['HS256'])
        user_id = payload.get('sub') or payload.get('user_id')
        if not user_id:
            raise HTTPException(status_code=401, detail="Invalid token")
        
        user = await db.users.find_one({"id": user_id}, {"_id": 0})
        if not user:
            raise HTTPException(status_code=401, detail="User not found")
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")


# ============== Pydantic Models ==============

class TaskPreview(BaseModel):
    id: str
    title: str
    type: str  # design, frontend, backend, qa
    description: str
    story_points: Optional[int] = None


class UserStoryPreview(BaseModel):
    id: str
    title: str
    description: str
    acceptance_criteria: List[str] = []
    story_points: Optional[int] = None
    tasks: List[TaskPreview] = []


class EpicPreview(BaseModel):
    id: str
    title: str
    description: str
    user_stories: List[UserStoryPreview] = []


class ParsedFeatureDocument(BaseModel):
    feature_title: str
    feature_description: str
    epics: List[EpicPreview] = []
    images: List[Dict[str, str]] = []  # base64 encoded images from document


class CreateArtifactsRequest(BaseModel):
    project_id: Optional[str] = None  # If None, create new project
    new_project_name: Optional[str] = None
    new_project_description: Optional[str] = None
    parsed_data: ParsedFeatureDocument


# ============== Helper Functions ==============

def extract_images_from_docx(doc: Document) -> List[Dict[str, str]]:
    """Extract all images from a DOCX document"""
    images = []
    
    try:
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                try:
                    image_data = rel.target_part.blob
                    # Determine image type from content type
                    content_type = rel.target_part.content_type
                    if "png" in content_type:
                        img_type = "png"
                    elif "jpeg" in content_type or "jpg" in content_type:
                        img_type = "jpeg"
                    elif "gif" in content_type:
                        img_type = "gif"
                    else:
                        img_type = "png"  # default
                    
                    # Convert to base64
                    base64_data = base64.b64encode(image_data).decode('utf-8')
                    images.append({
                        "data": f"data:image/{img_type};base64,{base64_data}",
                        "type": img_type,
                        "name": rel.target_ref.split('/')[-1] if '/' in rel.target_ref else rel.target_ref
                    })
                except Exception as e:
                    logger.warning(f"Failed to extract image: {e}")
    except Exception as e:
        logger.warning(f"Error accessing document relationships: {e}")
    
    return images


def extract_text_from_docx(doc: Document) -> str:
    """Extract all text from a DOCX document"""
    full_text = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            # Check if it's a heading
            if para.style and para.style.name and 'Heading' in para.style.name:
                full_text.append(f"\n## {text}\n")
            else:
                full_text.append(text)
    
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)
            if row_text:
                full_text.append(" | ".join(row_text))
    
    return "\n".join(full_text)


async def parse_with_ai(document_text: str) -> Dict[str, Any]:
    """Use AI to parse the document and generate structured artifacts"""
    from emergentintegrations.llm.chat import LlmChat, UserMessage
    from dotenv import load_dotenv
    
    load_dotenv()
    
    api_key = os.environ.get("EMERGENT_LLM_KEY")
    if not api_key:
        raise HTTPException(status_code=500, detail="AI parsing not configured. Missing EMERGENT_LLM_KEY.")
    
    system_prompt = """You are a Project Management Automation Engine.
Your task is to parse a Feature Document and generate structured project artifacts.

IMPORTANT: You must respond with ONLY valid JSON, no markdown code blocks, no explanation.

Output Format (strict JSON):
{
    "feature_title": "string - main feature name",
    "feature_description": "string - brief overview",
    "epics": [
        {
            "title": "string - Epic name in format: [Feature] - [Module/Flow]",
            "description": "string - Epic description",
            "user_stories": [
                {
                    "title": "string - User story title",
                    "description": "string - As a [user], I want to [action], so that [benefit]",
                    "acceptance_criteria": ["string - criterion 1", "string - criterion 2"],
                    "story_points": number or null,
                    "tasks": [
                        {"title": "string - [Story Title] - Design", "type": "design", "description": "UI/UX design task"},
                        {"title": "string - [Story Title] - Frontend", "type": "frontend", "description": "Frontend implementation"},
                        {"title": "string - [Story Title] - Backend", "type": "backend", "description": "Backend implementation"},
                        {"title": "string - [Story Title] - QA", "type": "qa", "description": "Testing and QA"}
                    ]
                }
            ]
        }
    ]
}

Rules:
1. Create 1+ Epics based on logical modules or functional groupings
2. Each Epic should have multiple User Stories
3. User Stories must follow format: "As a [user], I want to [action], so that [benefit]"
4. Each User Story generates 4 tasks: Design, Frontend, Backend, QA
5. Assign story points using Fibonacci scale (1,2,3,5,8) based on complexity
6. Extract acceptance criteria from the document"""

    chat = LlmChat(
        api_key=api_key,
        session_id=f"feature-parser-{uuid.uuid4()}",
        system_message=system_prompt
    ).with_model("openai", "gpt-4o")
    
    user_message = UserMessage(
        text=f"Parse this feature document and generate project artifacts:\n\n{document_text}"
    )
    
    try:
        response = await chat.send_message(user_message)
        
        # Clean up response - remove markdown code blocks if present
        response_text = response.strip()
        if response_text.startswith("```json"):
            response_text = response_text[7:]
        if response_text.startswith("```"):
            response_text = response_text[3:]
        if response_text.endswith("```"):
            response_text = response_text[:-3]
        response_text = response_text.strip()
        
        import json
        parsed = json.loads(response_text)
        return parsed
    except json.JSONDecodeError as e:
        logger.error(f"Failed to parse AI response as JSON: {e}")
        logger.error(f"Response was: {response[:500] if response else 'empty'}")
        raise HTTPException(status_code=500, detail="AI returned invalid JSON. Please try again.")
    except Exception as e:
        logger.error(f"AI parsing failed: {e}")
        raise HTTPException(status_code=500, detail=f"AI parsing failed: {str(e)}")


def add_ids_to_parsed_data(parsed: Dict[str, Any]) -> ParsedFeatureDocument:
    """Add unique IDs to all parsed artifacts for frontend tracking"""
    epics = []
    
    for epic_data in parsed.get("epics", []):
        user_stories = []
        
        for story_data in epic_data.get("user_stories", []):
            tasks = []
            
            for task_data in story_data.get("tasks", []):
                tasks.append(TaskPreview(
                    id=str(uuid.uuid4()),
                    title=task_data.get("title", ""),
                    type=task_data.get("type", "backend"),
                    description=task_data.get("description", ""),
                    story_points=task_data.get("story_points")
                ))
            
            user_stories.append(UserStoryPreview(
                id=str(uuid.uuid4()),
                title=story_data.get("title", ""),
                description=story_data.get("description", ""),
                acceptance_criteria=story_data.get("acceptance_criteria", []),
                story_points=story_data.get("story_points"),
                tasks=tasks
            ))
        
        epics.append(EpicPreview(
            id=str(uuid.uuid4()),
            title=epic_data.get("title", ""),
            description=epic_data.get("description", ""),
            user_stories=user_stories
        ))
    
    return ParsedFeatureDocument(
        feature_title=parsed.get("feature_title", "Untitled Feature"),
        feature_description=parsed.get("feature_description", ""),
        epics=epics,
        images=[]
    )


# ============== API Endpoints ==============

class ParseTextRequest(BaseModel):
    """Request model for parsing text content directly"""
    content: str
    title: Optional[str] = None


@router.post("/parse-text")
async def parse_text_content(
    request: ParseTextRequest,
    current_user: dict = Depends(get_current_user)
):
    """
    Parse text content directly (from Knowledge Base page) and return structured preview.
    """
    if not request.content or not request.content.strip():
        raise HTTPException(status_code=400, detail="Content is empty")
    
    try:
        document_text = request.content
        if request.title:
            document_text = f"# {request.title}\n\n{document_text}"
        
        logger.info(f"Parsing {len(document_text)} chars of text content")
        
        # Parse with AI
        parsed_data = await parse_with_ai(document_text)
        
        # Add IDs and structure the response
        result = add_ids_to_parsed_data(parsed_data)
        
        return {
            "success": True,
            "data": result.model_dump(),
            "document_text": document_text[:2000] + "..." if len(document_text) > 2000 else document_text
        }
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Failed to parse text content: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to parse content: {str(e)}")


@router.post("/parse")
async def parse_feature_document(
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user)
):
    """
    Parse a DOCX feature document and return structured preview of artifacts.
    Extracts text, images, and uses AI to generate Epics, User Stories, and Tasks.
    """
    # Validate file type
    if not file.filename.endswith('.docx'):
        raise HTTPException(status_code=400, detail="Only DOCX files are supported")
    
    try:
        # Read file content
        content = await file.read()
        doc = Document(BytesIO(content))
        
        # Extract text and images
        document_text = extract_text_from_docx(doc)
        images = extract_images_from_docx(doc)
        
        if not document_text.strip():
            raise HTTPException(status_code=400, detail="Document appears to be empty or unreadable")
        
        logger.info(f"Extracted {len(document_text)} chars and {len(images)} images from document")
        
        # Parse with AI
        parsed_data = await parse_with_ai(document_text)
        
        # Add IDs and structure the response
        result = add_ids_to_parsed_data(parsed_data)
        result.images = images
        
        return {
            "success": True,
            "data": result.model_dump(),
            "document_text": document_text[:2000] + "..." if len(document_text) > 2000 else document_text
        }
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Failed to parse document: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to parse document: {str(e)}")


@router.post("/create-artifacts")
async def create_artifacts(
    request: CreateArtifactsRequest,
    current_user: dict = Depends(get_current_user)
):
    """
    Create actual Epics, User Stories, and Tasks from parsed data.
    Can either link to existing project or create a new one.
    """
    now = datetime.now(timezone.utc).isoformat()
    
    try:
        # Determine project
        project_id = request.project_id
        
        if not project_id and request.new_project_name:
            # Create new project
            project_id = str(uuid.uuid4())
            new_project = {
                "id": project_id,
                "name": request.new_project_name,
                "description": request.new_project_description or request.parsed_data.feature_description,
                "project_type": "development",
                "status": "active",
                "created_by": current_user.get("id"),
                "created_at": now,
                "updated_at": now,
                "team_members": [current_user.get("id")],
                "tags": ["auto-generated", "feature-parser"]
            }
            await db.pm_projects.insert_one(new_project)
            logger.info(f"Created new project: {project_id}")
        
        if not project_id:
            raise HTTPException(status_code=400, detail="Either project_id or new_project_name is required")
        
        created_epics = []
        created_stories = []
        created_tasks = []
        
        # Create Epics
        for epic in request.parsed_data.epics:
            epic_id = str(uuid.uuid4())
            epic_doc = {
                "id": epic_id,
                "project_id": project_id,
                "title": epic.title,
                "description": epic.description,
                "status": "open",
                "priority": "medium",
                "created_by": current_user.get("id"),
                "created_at": now,
                "updated_at": now,
                "tags": ["auto-generated"]
            }
            await db.pm_epics.insert_one(epic_doc)
            created_epics.append(epic_id)
            
            # Create User Stories for this Epic
            for story in epic.user_stories:
                story_id = str(uuid.uuid4())
                story_doc = {
                    "id": story_id,
                    "project_id": project_id,
                    "epic_id": epic_id,
                    "name": story.title,  # Use 'name' for compatibility with TaskResponse model
                    "title": story.title,
                    "description": story.description,
                    "acceptance_criteria": "\n".join(story.acceptance_criteria) if story.acceptance_criteria else "",  # Convert list to string
                    "story_points": story.story_points,
                    "status": "draft",  # Use 'draft' for unassigned tasks - shown in Kanban 'Draft' column
                    "priority": "medium",
                    "type": "user_story",
                    "parent_task_id": None,  # Explicitly set for top-level tasks
                    "created_by": current_user.get("id"),
                    "created_at": now,
                    "updated_at": now,
                    "tags": ["auto-generated"]
                }
                await db.pm_tasks.insert_one(story_doc)
                created_stories.append(story_id)
                
                # Create Tasks for this User Story
                for task in story.tasks:
                    task_id = str(uuid.uuid4())
                    task_doc = {
                        "id": task_id,
                        "project_id": project_id,
                        "epic_id": epic_id,
                        "parent_task_id": story_id,  # Link to user story (subtask)
                        "name": task.title,  # Use 'name' for compatibility with TaskResponse model
                        "title": task.title,
                        "description": task.description,
                        "story_points": task.story_points,
                        "status": "draft",  # Use 'draft' for unassigned tasks - shown in Kanban 'Draft' column
                        "priority": "medium",
                        "type": task.type,  # design, frontend, backend, qa
                        "created_by": current_user.get("id"),
                        "created_at": now,
                        "updated_at": now,
                        "tags": ["auto-generated", task.type]
                    }
                    await db.pm_tasks.insert_one(task_doc)
                    created_tasks.append(task_id)
        
        return {
            "success": True,
            "message": f"Successfully created {len(created_epics)} epics, {len(created_stories)} user stories, and {len(created_tasks)} tasks",
            "project_id": project_id,
            "created": {
                "epics": len(created_epics),
                "user_stories": len(created_stories),
                "tasks": len(created_tasks)
            }
        }
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Failed to create artifacts: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to create artifacts: {str(e)}")


@router.get("/projects")
async def get_engineering_projects(
    current_user: dict = Depends(get_current_user)
):
    """Get list of engineering projects for selection"""
    
    projects = await db.pm_projects.find(
        {"project_type": "development"},
        {"_id": 0, "id": 1, "name": 1, "description": 1, "status": 1}
    ).to_list(100)
    
    return projects


# Register the router
def register_routes(app):
    app.include_router(router, prefix="/api/engineering")
    logger.info("Feature Parser routes loaded successfully")
